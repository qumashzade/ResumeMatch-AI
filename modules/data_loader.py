"""
data_loader.py — CV və iş elanı fayllarını diskdən yükləyir.

Dəstəklənən formatlar: .txt, .pdf, .docx, .jpg, .jpeg, .png
"""

import os
import glob
import logging
from dataclasses import dataclass, field
from typing import List, Optional

logger = logging.getLogger(__name__)


@dataclass
class Document:
    """Yüklənmiş sənəd (CV və ya iş elanı)."""
    filename: str
    raw_text: str
    filepath: str
    file_type: str
    metadata: dict = field(default_factory=dict)


def _extract_text_from_pdf(filepath: str) -> str:
    """PDF fayldan mətn çıxarır."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(filepath)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)
    except ImportError:
        logger.warning("PyPDF2 quraşdırılmayıb. pip install PyPDF2")
        return ""
    except Exception as e:
        logger.error("PDF oxuma xətası %s: %s", filepath, e)
        return ""


def _extract_text_from_docx(filepath: str) -> str:
    """DOCX fayldan mətn çıxarır."""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except ImportError:
        logger.warning("python-docx quraşdırılmayıb. pip install python-docx")
        return ""
    except Exception as e:
        logger.error("DOCX oxuma xətası %s: %s", filepath, e)
        return ""


def _extract_text_from_txt(filepath: str) -> str:
    """Mətn faylını oxuyur."""
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            with open(filepath, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    logger.error("Fayl decode edilə bilmədi: %s", filepath)
    return ""


def _extract_text_from_image(filepath: str) -> str:
    """OCR ilə şəkildən mətn çıxarır."""
    try:
        import pytesseract
        from PIL import Image, ImageEnhance

        # Windows üçün Tesseract yolunu yoxla
        import shutil
        if not shutil.which("tesseract"):
            common_paths = [
                r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
            ]
            for p in common_paths:
                if os.path.isfile(p):
                    pytesseract.pytesseract.tesseract_cmd = p
                    break

        img = Image.open(filepath)
        img = img.convert("L")
        img = ImageEnhance.Contrast(img).enhance(2.0)
        img = ImageEnhance.Sharpness(img).enhance(2.0)
        w, h = img.size
        if w < 1500:
            scale = 1500 / w
            img = img.resize((int(w * scale), int(h * scale)), Image.LANCZOS)

        text = pytesseract.image_to_string(img, lang="eng")
        logger.info("OCR: %d simvol çıxarıldı — %s", len(text), filepath)
        return text

    except ImportError:
        logger.warning("pytesseract/Pillow quraşdırılmayıb. pip install pytesseract Pillow")
        return ""
    except Exception as e:
        logger.error("OCR xətası %s: %s", filepath, e)
        return ""


def load_document(filepath: str) -> Optional[Document]:
    """Bir sənədi diskdən yükləyir."""
    if not os.path.isfile(filepath):
        logger.warning("Fayl tapılmadı: %s", filepath)
        return None

    filename = os.path.basename(filepath)
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        text = _extract_text_from_pdf(filepath)
        file_type = "pdf"
    elif ext in (".docx", ".doc"):
        text = _extract_text_from_docx(filepath)
        file_type = "docx"
    elif ext in (".txt", ".text", ".md"):
        text = _extract_text_from_txt(filepath)
        file_type = "txt"
    elif ext in (".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".webp"):
        text = _extract_text_from_image(filepath)
        file_type = "image"
    else:
        logger.warning("Dəstəklənməyən format: %s", filepath)
        return None

    if not text.strip():
        logger.warning("Boş məzmun: %s", filepath)
        return None

    return Document(filename=filename, raw_text=text, filepath=filepath, file_type=file_type)


def load_cvs_from_folder(folder_path: str) -> List[Document]:
    """Qovluqdakı bütün CV fayllarını yükləyir."""
    if not os.path.isdir(folder_path):
        raise FileNotFoundError(f"Qovluq tapılmadı: {folder_path}")

    patterns = ["*.txt", "*.pdf", "*.docx", "*.doc", "*.jpg", "*.jpeg", "*.png", "*.bmp", "*.tiff", "*.webp"]
    files = []
    for pattern in patterns:
        files.extend(glob.glob(os.path.join(folder_path, pattern)))

    files = sorted(set(files))
    logger.info("%d fayl tapıldı: %s", len(files), folder_path)

    documents = []
    for fp in files:
        doc = load_document(fp)
        if doc:
            documents.append(doc)

    logger.info("%d CV uğurla yükləndi", len(documents))
    return documents


def load_job_description(filepath: str) -> Document:
    """İş elanı faylını yükləyir."""
    doc = load_document(filepath)
    if doc is None:
        raise ValueError(f"İş elanı yüklənə bilmədi: {filepath}")
    return doc
